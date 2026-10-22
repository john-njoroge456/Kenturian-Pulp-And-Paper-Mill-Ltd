from datetime import date, datetime, time, timedelta
from io import BytesIO

from django.db.models import Count, Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import OrderItem, OrderRequest

NAVY = "14335C"

ORDER_COLUMNS = [
    # (header, width)
    ("Reference", 13), ("Order date", 17), ("Company", 26), ("Contact person", 20),
    ("Phone", 16), ("Email", 26), ("Item #", 7), ("Item", 22), ("Colour", 14),
    ("Length (mm)", 12), ("GSM", 7), ("Quantity (kg)", 13), ("Status", 11),
    ("Preferred delivery", 17), ("Delivery location",
                                 22), ("LPO attached", 12), ("Notes", 40),
]


def _style_header(ws, columns):
    fill = PatternFill("solid", start_color=NAVY)
    for col, (title, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"


def orders_workbook(orders, product=""):
    """One row per item, so staff can sort, filter and pivot in Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Order items"
    _style_header(ws, ORDER_COLUMNS)

    row = 1
    for order in orders.prefetch_related("items").order_by("-created_at"):
        created = timezone.localtime(order.created_at).replace(tzinfo=None)
        for item in order.items.all():
            if product and item.item_description != product:
                continue
            ws.append([
                order.reference, created, order.company_name, order.contact_person,
                order.phone, order.email, item.position, item.item_description,
                item.color, float(item.length), item.gsm, item.quantity,
                order.get_status_display(), order.preferred_delivery_date,
                order.delivery_location, "Yes" if order.lpo_file else "No", order.notes,
            ])
            row += 1

    for r in range(2, row + 1):
        ws.cell(row=r, column=2).number_format = "dd mmm yyyy hh:mm"
        ws.cell(row=r, column=12).number_format = "#,##0"
        ws.cell(row=r, column=14).number_format = "dd mmm yyyy"
        ws.cell(row=r, column=17).alignment = Alignment(
            wrap_text=True, vertical="top")

    if row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(ORDER_COLUMNS))}{row}"
        total_row = row + 2
        ws.cell(row=total_row, column=11, value="Total").font = Font(bold=True)
        total = ws.cell(row=total_row, column=12, value=f"=SUM(L2:L{row})")
        total.font = Font(bold=True)
        total.number_format = "#,##0"
    return wb


def xlsx_response(wb, filename):
    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

# ---------------------------------------------------------------------------
# Order report (admin "Order report" page)
# ---------------------------------------------------------------------------


PRESETS = [
    ("this_month", "This month"),
    ("last_month", "Last month"),
    ("last_30", "Last 30 days"),
    ("this_year", "This year"),
]


def _parse_date(value):
    try:
        return date.fromisoformat(value) if value else None
    except ValueError:
        return None


def resolve_period(params):
    """Work out the report dates from ?preset= or ?start=&end=."""
    today = timezone.localdate()
    start, end = _parse_date(params.get(
        "start")), _parse_date(params.get("end"))
    if start and end:
        if start > end:
            start, end = end, start
        return "custom", start, end

    preset = params.get("preset") or "this_month"
    if preset == "last_month":
        end = today.replace(day=1) - timedelta(days=1)
        start = end.replace(day=1)
    elif preset == "last_30":
        start, end = today - timedelta(days=29), today
    elif preset == "this_year":
        start, end = today.replace(month=1, day=1), today
    else:
        preset = "this_month"
        start, end = today.replace(day=1), today
    return preset, start, end


def _change(current, previous):
    if not previous:
        return None
    return round((current - previous) * 100 / previous)


def _pct(part, whole):
    return round(part * 100 / whole) if whole else 0


def _scope(orders, client, product, status):
    if client:
        orders = orders.filter(company_name__iexact=client)
    if status:
        orders = orders.filter(status=status)
    if product:
        orders = orders.filter(items__item_description=product)
    # Re-select by id so joins never double-count orders
    return OrderRequest.objects.filter(pk__in=orders.values("pk"))


def _scoped_items(orders, product):
    items = OrderItem.objects.filter(order__in=orders.values("pk"))
    return items.filter(item_description=product) if product else items


def _between(start, end):
    """Whole days from start to end (Nairobi time), without asking the database to convert time zones."""
    tz = timezone.get_current_timezone()
    return {
        "created_at__gte": timezone.make_aware(datetime.combine(start, time.min), tz),
        "created_at__lt": timezone.make_aware(datetime.combine(end + timedelta(days=1), time.min), tz),
    }


def build_order_report(start, end, client="", product="", status=""):
    days = (end - start).days + 1
    prev_end = start - timedelta(days=1)
    prev_start = prev_end - timedelta(days=days - 1)

    orders = _scope(OrderRequest.objects.filter(
        **_between(start, end)), client, product, status)
    prev_orders = _scope(OrderRequest.objects.filter(
        **_between(prev_start, prev_end)), client, product, status)
    items = _scoped_items(orders, product)
    prev_items = _scoped_items(prev_orders, product)

    order_count = orders.count()
    total_kg = items.aggregate(t=Sum("quantity"))["t"] or 0
    lines = items.count()
    prev_count = prev_orders.count()
    prev_kg = prev_items.aggregate(t=Sum("quantity"))["t"] or 0

    # Products, each with its colours underneath
    by_product = list(
        items.values("item_description")
        .annotate(kg=Sum("quantity"), lines=Count("id"), orders=Count("order", distinct=True))
        .order_by("-kg"))
    colours = {}
    for r in (items.values("item_description", "color")
              .annotate(kg=Sum("quantity"), lines=Count("id"), orders=Count("order", distinct=True))
              .order_by("-kg")):
        colours.setdefault(r["item_description"], []).append(r)
    for p in by_product:
        p["share"] = _pct(p["kg"], total_kg)
        p["colours"] = colours.get(p["item_description"], [])
        for c in p["colours"]:
            c["share"] = _pct(c["kg"], total_kg)
            c["share_of_product"] = _pct(c["kg"], p["kg"])

    status_counts = {r["status"]: r["n"]
                     for r in orders.values("status").annotate(n=Count("id"))}
    status_kg = {r["order__status"]: r["kg"] for r in
                 items.values("order__status").annotate(kg=Sum("quantity"))}
    by_status = [
        {
            "code": code, "label": label,
            "orders": status_counts.get(code, 0),
            "kg": status_kg.get(code) or 0,
            "share": _pct(status_counts.get(code, 0), order_count),
        }
        for code, label in OrderRequest.STATUS_CHOICES
    ]

    top_clients = list(
        items.values("order__company_name")
        .annotate(kg=Sum("quantity"), orders=Count("order", distinct=True))
        .order_by("-kg")[:10])
    for r in top_clients:
        r["share"] = _pct(r["kg"], total_kg)

    stale_new = OrderRequest.objects.filter(
        status="new", created_at__lt=timezone.now() - timedelta(days=2)).count()

    status_label = dict(OrderRequest.STATUS_CHOICES).get(status, "")
    active = [f"Client: {client}" if client else "", f"Product: {product}" if product else "",
              f"Status: {status_label}" if status else ""]

    return {
        "start": start, "end": end, "days": days,
        "prev_start": prev_start, "prev_end": prev_end,
        "orders": orders,
        "order_count": order_count, "total_kg": total_kg, "lines": lines,
        "avg_kg": round(total_kg / order_count) if order_count else 0,
        "prev_count": prev_count, "prev_kg": prev_kg,
        "count_change": _change(order_count, prev_count),
        "kg_change": _change(total_kg, prev_kg),
        "by_product": by_product, "by_status": by_status, "top_clients": top_clients,
        "stale_new": stale_new,
        "client": client, "product": product, "status": status,
        "filters_text": " · ".join(a for a in active if a) or "All orders",
        "client_options": sorted(set(
            OrderRequest.objects.order_by().values_list("company_name", flat=True)), key=str.lower),
        "product_options": sorted(set(
            OrderItem.objects.order_by().values_list("item_description", flat=True)), key=str.lower),
        "status_options": OrderRequest.STATUS_CHOICES,
    }


def _table_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    _style_header(ws, columns)
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(columns) + 1):
            if isinstance(ws.cell(row=r, column=c).value, int):
                ws.cell(row=r, column=c).number_format = "#,##0"
    return ws


def report_workbook(report):
    wb = orders_workbook(report["orders"], product=report["product"])

    summary = wb.create_sheet("Summary", 0)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 18
    summary["A1"] = "Kenturian Pulp & Paper Mills Limited — Order report"
    summary["A1"].font = Font(bold=True, size=14, color=NAVY)
    summary["A2"] = f"{report['start']:%d %b %Y} to {report['end']:%d %b %Y}"
    summary["A3"] = f"Filters: {report['filters_text']}"
    summary.append([])
    summary.append(["", "This period", "Previous period"])
    for cell in summary[5]:
        cell.font = Font(bold=True)
    summary.append(["Orders", report["order_count"], report["prev_count"]])
    summary.append(
        ["Total quantity (kg)", report["total_kg"], report["prev_kg"]])
    summary.append(["Item lines", report["lines"], ""])
    summary.append(["Average kg per order", report["avg_kg"], ""])
    for r in range(6, 10):
        for c in (2, 3):
            summary.cell(row=r, column=c).number_format = "#,##0"

    # Products with their colours underneath
    ws = wb.create_sheet("By product & colour")
    _style_header(ws, [("Product", 26), ("Colour", 20), ("Quantity (kg)", 14),
                       ("Share of total %", 16), ("Item lines", 11), ("Orders", 9)])
    for p in report["by_product"]:
        ws.append([p["item_description"], "All colours",
                  p["kg"], p["share"], p["lines"], p["orders"]])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for c in p["colours"]:
            ws.append(["", c["color"] or "—", c["kg"],
                      c["share"], c["lines"], c["orders"]])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = "#,##0"

    if not report["status"]:
        _table_sheet(wb, "By status",
                     [("Status", 14), ("Orders", 9), ("Quantity (kg)", 14)],
                     [[r["label"], r["orders"], r["kg"]] for r in report["by_status"]])
    if not report["client"]:
        _table_sheet(wb, "Top clients",
                     [("Company", 30), ("Quantity (kg)", 14),
                      ("Share %", 9), ("Orders", 9)],
                     [[r["order__company_name"], r["kg"], r["share"], r["orders"]] for r in report["top_clients"]])
    return wb


APPLICATION_COLUMNS = [
    ("Reference", 13), ("Applied on", 17), ("Position",
                                            24), ("Name", 22), ("Email", 26),
    ("Phone", 16), ("Status", 14), ("Message",
                                    45), ("HR notes", 35), ("CV link", 40),
]


def applications_workbook(applications, request):
    from django.urls import reverse
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    _style_header(ws, APPLICATION_COLUMNS)
    row = 1
    for app in applications.select_related("job").order_by("-created_at"):
        cv_url = request.build_absolute_uri(
            reverse("admin:main_jobapplication_cv", args=[app.pk])) if app.cv_file else ""
        ws.append([
            app.reference, timezone.localtime(
                app.created_at).replace(tzinfo=None),
            app.job.title if app.job else "Speculative", app.name, app.email, app.phone,
            app.get_status_display(), app.message, app.hr_notes, cv_url,
        ])
        row += 1
    for r in range(2, row + 1):
        ws.cell(row=r, column=2).number_format = "dd mmm yyyy hh:mm"
        for c in (8, 9):
            ws.cell(row=r, column=c).alignment = Alignment(
                wrap_text=True, vertical="top")
    if row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(APPLICATION_COLUMNS))}{row}"
    return wb
